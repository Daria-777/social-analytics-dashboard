from datetime import datetime, timedelta, timezone
from io import BytesIO
from uuid import UUID
from openpyxl import load_workbook
from app.database import get_session
from app.models import AccountSnapshot, ContentSnapshot


def add_snapshots(client, content):
    gen = client.app.dependency_overrides[get_session]()
    db = next(gen)
    now = datetime.now(timezone.utc) - timedelta(seconds=5)
    try:
        db.add_all([
            AccountSnapshot(account_id=UUID(content['account_id']), source='tiktok_api', metric_scope='current', snapshot_at=now-timedelta(days=2), followers=10, snapshot_status='confirmed'),
            AccountSnapshot(account_id=UUID(content['account_id']), source='tiktok_api', metric_scope='current', snapshot_at=now, followers=9, snapshot_status='confirmed'),
            AccountSnapshot(account_id=UUID(content['account_id']), source='tiktok_studio', metric_scope='current', snapshot_at=now, followers=90, snapshot_status='manual', raw_payload={'agent_run_id':'private-run', 'access_token':'secret-export-marker'}),
            ContentSnapshot(content_id=UUID(content['id']), source='tiktok_api', metric_scope='lifetime', snapshot_at=now, views=0, likes=None, comments=0, snapshot_status='confirmed', raw_payload={'access_token':'secret-export-marker'}),
            ContentSnapshot(content_id=UUID(content['id']), source='tiktok_api', metric_scope='lifetime', snapshot_at=now-timedelta(days=40), views=999, snapshot_status='confirmed'),
        ])
        db.commit()
    finally: gen.close()


def test_summary_period_sources_and_negative_change(client, content):
    add_snapshots(client, content)
    result = client.get('/reports/summary?days=7&timezone=Europe/Moscow')
    assert result.status_code == 200
    summary = result.json(); group = summary['groups'][0]
    assert summary['timezone'] == 'Europe/Moscow'
    assert group['measurements'][0]['followers'] == 9
    assert group['measurements'][0]['change'] == -1
    assert len(group['measurements']) == 1
    assert '90 подписчиков' not in summary['text']
    assert group['tops'][0]['posts'][0]['views'] == 0
    assert '999' not in summary['text']
    assert 'secret-export-marker' not in result.text
    assert 'private-run' not in result.text


def test_excel_literal_strings_unknown_zero_no_secrets(client, content):
    add_snapshots(client, content)
    assert client.patch('/content/'+content['id']+'/tags', json={'internal_title':'=HYPERLINK("https://evil.invalid")'}).status_code == 200
    response = client.get('/reports/export.xlsx?days=7&timezone=Europe/Moscow')
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    assert '.xlsx"' in response.headers['content-disposition']
    workbook = load_workbook(BytesIO(response.content))
    assert len(workbook.sheetnames) == 6
    publications = workbook['Публикации']
    assert publications['D2'].data_type == 's'
    assert publications['D2'].value.startswith('=HYPERLINK')
    ws = workbook['Измерения публикаций']
    headers = [c.value for c in ws[1]]
    assert ws.cell(2, headers.index('Просмотры')+1).value == 0
    assert ws.cell(2, headers.index('Лайки')+1).value is None
    assert ws.cell(2, headers.index('Комментарии')+1).value == 0
    assert ws.max_row == 2
    account_ws = workbook['Измерения аккаунтов']
    assert any('AI-агент' in row for row in account_ws.values)
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert cell.data_type != 'f'
                assert 'secret-export-marker' not in str(cell.value)
                assert 'private-run' not in str(cell.value)


def test_exports_auth_validation_and_empty_pdf(client):
    for path in ['/reports/summary', '/reports/export.xlsx', '/reports/export.pdf']:
        assert client.get(path+'?days=8').status_code == 422
        assert client.get(path+'?timezone=Not/AZone').status_code == 422
    assert client.get('/reports/export.csv').status_code == 422
    pdf = client.get('/reports/export.pdf')
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF-')
    assert b'/FontFile2' in pdf.content
    client.headers.pop('X-API-Key')
    for path in ['/reports/summary','/reports/export.xlsx','/reports/export.pdf']:
        assert client.get(path).status_code == 401


def test_pdf_long_caption_and_literal_markup(client, content):
    add_snapshots(client, content)
    assert client.patch('/content/'+content['id']+'/tags', json={'internal_title':'<b>Привет</b> & ' + 'Очень длинное имя ' * 40}).status_code == 200
    response = client.get('/reports/export.pdf?days=30')
    assert response.status_code == 200 and response.content.startswith(b'%PDF-')
    assert b'secret-export-marker' not in response.content


def test_local_calendar_boundary_and_limit(client, content, monkeypatch):
    from app import reports
    gen=client.app.dependency_overrides[get_session](); db=next(gen)
    try:
        report = reports.load_report(db, 7, 'Europe/Moscow', now=datetime(2026, 10, 6, 22, tzinfo=timezone.utc))
        assert report.start == datetime(2026, 9, 30, 21, tzinfo=timezone.utc)
    finally: gen.close()
    monkeypatch.setattr(reports, 'MAX_ROWS', 0)
    assert client.get('/reports/export.xlsx').status_code == 413


def test_summary_falls_back_to_ui_without_api(client, content):
    gen=client.app.dependency_overrides[get_session](); db=next(gen)
    try:
        db.add(AccountSnapshot(account_id=UUID(content['account_id']), source='tiktok_studio', metric_scope='current', snapshot_at=datetime.now(timezone.utc)-timedelta(seconds=1), followers=0, snapshot_status='manual'))
        db.commit()
    finally: gen.close()
    result=client.get('/reports/summary').json()['groups'][0]['measurements']
    assert result[0]['source'] == 'TikTok Studio'
    assert result[0]['followers'] == 0
    assert result[0]['change'] is None
