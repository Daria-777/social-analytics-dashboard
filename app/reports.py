"""Read-only exports from saved measurements. No provider calls or credentials."""
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, time, timedelta, timezone
from decimal import Decimal
from enum import Enum
from io import BytesIO
from pathlib import Path
from threading import Lock
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models as m
from app.analytics_api import collection_method
from app.enums import DataStatus

MAX_ROWS = 10000
SOURCE = {'instagram_api': 'Instagram API', 'instagram_ui': 'Instagram Insights', 'tiktok_api': 'TikTok API', 'tiktok_studio': 'TikTok Studio', 'manual': 'Вручную'}
METHOD = {'automatic': 'API', 'agent': 'AI-агент', 'manual': 'Вручную', 'unknown': 'Не указан'}
SAFE_STATUS = {DataStatus.CONFIRMED, DataStatus.MANUAL}


@dataclass
class Report:
    start: datetime
    end: datetime
    zone: ZoneInfo
    accounts: list
    content: list
    account_rows: list
    content_rows: list
    retention: list
    geography: list

    @property
    def period(self):
        return f'{self.start.astimezone(self.zone):%d.%m.%Y} - {self.end.astimezone(self.zone):%d.%m.%Y}'


def bounded(db, statement):
    rows = list(db.scalars(statement.limit(MAX_ROWS + 1)))
    if len(rows) > MAX_ROWS:
        raise HTTPException(413, 'Слишком много измерений. Выберите более короткий период')
    return rows


def load_report(db: Session, days: int, zone_name: str, now=None):
    try:
        zone = ZoneInfo(zone_name)
    except (ZoneInfoNotFoundError, ValueError):
        raise HTTPException(422, 'Неизвестный часовой пояс') from None
    end = now or datetime.now(timezone.utc)
    start_day = end.astimezone(zone).date() - timedelta(days=days - 1)
    start = datetime.combine(start_day, time.min, zone).astimezone(timezone.utc)
    accounts = bounded(db, select(m.SocialAccount).order_by(m.SocialAccount.platform, m.SocialAccount.username, m.SocialAccount.id))
    content = bounded(db, select(m.Content).order_by(m.Content.published_at, m.Content.id))
    def measurements(model):
        return bounded(db, select(model).where(model.snapshot_at >= start, model.snapshot_at <= end).order_by(model.snapshot_at, model.created_at, model.id))
    account_rows = measurements(m.AccountSnapshot)
    content_rows = measurements(m.ContentSnapshot)
    retention = measurements(m.ContentRetentionSnapshot)
    geography = bounded(db, select(m.ContentGeoSnapshot).join(m.ContentSnapshot, m.ContentGeoSnapshot.content_snapshot_id == m.ContentSnapshot.id).where(m.ContentSnapshot.snapshot_at >= start, m.ContentSnapshot.snapshot_at <= end).order_by(m.ContentGeoSnapshot.content_snapshot_id, m.ContentGeoSnapshot.country_code))
    referenced = {s.content_id for s in content_rows + retention}
    content = [c for c in content if c.id in referenced or c.published_at and start <= c.published_at <= end]
    return Report(start, end, zone, accounts, content, account_rows, content_rows, retention, geography)


def number(value):
    return 'нет данных' if value is None else f'{value:,}'.replace(',', ' ')


def title(content):
    return (content.internal_title or content.caption or 'Без названия').strip()[:160]


def report_summary(report):
    """No addition across sources, no invented weekly totals from lifetime counts."""
    groups = []
    posts = {c.id: c for c in report.content}
    for account in report.accounts:
        source_rows = defaultdict(list)
        for row in report.account_rows:
            if row.account_id == account.id and row.followers is not None and row.metric_scope in {'current', 'lifetime'} and row.source_period_start is None and row.source_period_end is None and row.snapshot_status in SAFE_STATUS:
                source_rows[row.source].append(row)
        top_rows = defaultdict(dict)
        for row in report.content_rows:
            post = posts.get(row.content_id)
            if post and post.account_id == account.id and row.metric_scope == 'lifetime' and row.source_period_start is None and row.source_period_end is None and row.snapshot_status in SAFE_STATUS and row.views is not None:
                top_rows[row.source][row.content_id] = row
        measurements = []
        for source, rows in source_rows.items():
            first, last = rows[0], rows[-1]
            measurements.append({'source': SOURCE[source], 'followers': last.followers, 'at': last.snapshot_at.astimezone(report.zone).strftime('%d.%m %H:%M'), 'change': last.followers - first.followers if len(rows) > 1 and first.snapshot_at < last.snapshot_at else None, 'between': f'{first.snapshot_at.astimezone(report.zone):%d.%m %H:%M} - {last.snapshot_at.astimezone(report.zone):%d.%m %H:%M}'})
        tops = []
        for source, rows in top_rows.items():
            best = sorted(rows.values(), key=lambda s: (-s.views, str(s.content_id)))[:5]
            tops.append({'source': SOURCE[source], 'posts': [{'title': title(posts[s.content_id]), 'views': s.views, 'likes': s.likes, 'comments': s.comments, 'shares': s.shares, 'at': s.snapshot_at.astimezone(report.zone).strftime('%d.%m %H:%M')} for s in best]})
        # A brief report selects one source per metric, never sums providers/UI.
        preferred = 'instagram_api' if account.platform == 'instagram' else 'tiktok_api'
        rank = lambda label: 0 if label == SOURCE[preferred] else 1 if label in ('Instagram Insights', 'TikTok Studio') else 2
        measurements = sorted(measurements, key=lambda item: rank(item['source']))[:1]
        tops = sorted(tops, key=lambda item: rank(item['source']))[:1]
        groups.append({'platform': 'Instagram' if account.platform == 'instagram' else 'TikTok', 'username': account.username[:80], 'measurements': measurements, 'tops': tops, 'new_posts': sum(c.account_id == account.id and c.published_at is not None and report.start <= c.published_at <= report.end for c in report.content)})
    lines = [f'DASH - отчёт {report.period}', f'Часовой пояс: {report.zone.key}']
    for group in groups:
        lines += ['', f'{group["platform"]} @{group["username"]}', f'Новых публикаций: {group["new_posts"]}']
        if not group['measurements']:
            lines.append('Подписчики: нет измерений за период')
        for item in group['measurements']:
            lines.append(f'{item["source"]}: {number(item["followers"])} подписчиков ({item["at"]})')
            if item['change'] is None:
                lines.append('Изменение: нужен второй замер')
            else:
                lines.append(f'Изменение между замерами: {item["change"]:+d} ({item["between"]})')
        for block in group['tops']:
            lines.append(f'Просмотры публикаций, накопленные - {block["source"]}:')
            for post in block['posts'][:3]:
                lines.append(f'• {post["title"][:70]}: {number(post["views"])}')
    lines += ['', 'Сводка предпочитает API; дополнительные источники - в Excel. Период выбирает даты измерений. Накопленные просмотры не равны просмотрам за неделю. Полные таблицы - в Excel.']
    text = '\n'.join(lines)
    if len(text) > 3500:
        text = text[:3370].rsplit('\n', 1)[0] + '\n\nСводка сокращена. Полные данные доступны в Excel и PDF.'
    return {'period': report.period, 'timezone': report.zone.key, 'groups': groups, 'text': text}


# Explicit field lists exclude raw payloads, credentials, signed media URLs and run logs.
COMMON = [('snapshot_at', 'Измерение (UTC)'), ('source', 'Источник'), ('metric_scope', 'Область метрики'), ('source_period_start', 'Начало периода (UTC)'), ('source_period_end', 'Конец периода (UTC)'), ('snapshot_status', 'Статус')]
ACCOUNT_METRICS = [('followers', 'Подписчики'), ('views', 'Просмотры'), ('reach', 'Охват'), ('unique_viewers', 'Уникальные зрители'), ('profile_views', 'Посещения профиля'), ('likes', 'Лайки'), ('comments', 'Комментарии'), ('shares', 'Репосты'), ('saves', 'Сохранения'), ('new_viewers', 'Новые зрители')]
CONTENT_METRICS = [('views', 'Просмотры'), ('reach', 'Охват'), ('unique_viewers', 'Уникальные зрители'), ('likes', 'Лайки'), ('comments', 'Комментарии'), ('shares', 'Репосты'), ('saves', 'Сохранения'), ('followers_gained', 'Новые подписчики'), ('profile_visits', 'Посещения профиля'), ('profile_activity', 'Действия в профиле'), ('watch_time_total_seconds', 'Общее время просмотра (с)'), ('watch_time_avg_seconds', 'Среднее время просмотра (с)'), ('completion_rate', 'Досмотры (%)'), ('traffic_for_you_pct', 'Рекомендации (%)'), ('traffic_profile_pct', 'Профиль (%)'), ('traffic_search_pct', 'Поиск (%)'), ('traffic_following_pct', 'Лента подписок (%)'), ('traffic_other_pct', 'Другой трафик (%)'), ('new_viewers_pct', 'Новые зрители (%)'), ('returning_viewers_pct', 'Вернувшиеся зрители (%)'), ('female_pct', 'Женщины (%)'), ('male_pct', 'Мужчины (%)'), ('other_gender_pct', 'Другой пол (%)'), ('age_18_24_pct', '18–24 (%)'), ('age_25_34_pct', '25–34 (%)'), ('age_35_44_pct', '35–44 (%)'), ('age_45_54_pct', '45–54 (%)'), ('age_55_plus_pct', '55+ (%)')]
CONTENT_FIELDS = [('published_at', 'Опубликовано (UTC)'), ('content_type', 'Тип'), ('duration_seconds', 'Длительность (с)'), ('caption', 'Подпись'), ('hook_text', 'Хук'), ('hook_type', 'Тип хука'), ('topic', 'Тема'), ('series', 'Серия'), ('content_pillar', 'Направление'), ('format', 'Формат'), ('cta_type', 'CTA')]


def xlsx_report(report):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(write_only=True)
    accounts = {a.id: a for a in report.accounts}
    posts = {c.id: c for c in report.content}

    def values(obj, fields):
        return [getattr(obj, key, None) for key, _ in fields]

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.freeze_panes = 'A2'
        from openpyxl.utils import get_column_letter
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(42, max(18, len(header) + 3))
        def cells(data, header=False):
            result = []
            for value in data:
                if isinstance(value, Enum): value = value.value
                if isinstance(value, datetime): value = value.astimezone(timezone.utc).replace(tzinfo=None)
                if isinstance(value, Decimal): value = float(value)
                if isinstance(value, int) and abs(value) > 999999999999999: value = str(value)
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub('', value)
                    if len(value) > 32767:
                        raise HTTPException(413, 'Текст слишком длинный для Excel')
                cell = WriteOnlyCell(ws, value=value)
                # Force all text to literal strings, including Excel formula prefixes.
                if isinstance(value, str): cell.data_type = 's'
                if isinstance(value, datetime): cell.number_format = 'yyyy-mm-dd hh:mm:ss'
                if header:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='067462')
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                result.append(cell)
            return result
        ws.append(cells(headers, True))
        count = 1
        for row in rows:
            ws.append(cells(row)); count += 1
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{count}'

    sheet('Об отчёте', ['Параметр', 'Значение'], [
        ['Период измерений', report.period], ['Часовой пояс периода', report.zone.key], ['Сформирован (UTC)', report.end],
        ['Даты в таблицах', 'UTC'], ['Пустая ячейка', 'Нет данных; это не ноль'],
        ['lifetime', 'Накопленный счётчик, не результат за выбранную неделю'], ['range', 'Показатель за указанный в строке период'],
        ['Источники', 'Данные API и интерфейса не суммируются'], ['Изменение подписчиков', 'Разница между первым и последним доступными замерами периода'],
    ])
    sheet('Публикации', ['ID', 'Платформа', 'Аккаунт', 'Название'] + [h for _, h in CONTENT_FIELDS], ([str(c.id), c.platform, accounts[c.account_id].username, title(c)] + values(c, CONTENT_FIELDS) for c in report.content))
    sheet('Измерения аккаунтов', ['Платформа', 'Аккаунт', 'Способ сбора'] + [h for _, h in COMMON + ACCOUNT_METRICS], ([accounts[s.account_id].platform, accounts[s.account_id].username, METHOD[collection_method(s)]] + values(s, COMMON + ACCOUNT_METRICS) for s in report.account_rows))
    sheet('Измерения публикаций', ['ID публикации', 'Платформа', 'Название', 'Способ сбора'] + [h for _, h in COMMON + CONTENT_METRICS], ([str(s.content_id), posts[s.content_id].platform, title(posts[s.content_id]), METHOD[collection_method(s)]] + values(s, COMMON + CONTENT_METRICS) for s in report.content_rows))
    retention_fields = [(k, h) for k, h in COMMON if k != 'metric_scope'] + [('drop_off_second', 'Отвал (с)'), ('average_watch_seconds', 'Среднее время (с)'), ('completion_rate', 'Досмотры (%)')]
    sheet('Удержание', ['ID публикации', 'Название', 'Способ сбора'] + [h for _, h in retention_fields], ([str(s.content_id), title(posts[s.content_id]), METHOD[collection_method(s)]] + values(s, retention_fields) for s in report.retention))
    snapshots = {s.id: s for s in report.content_rows}
    sheet('География', ['ID публикации', 'Название', 'Измерение (UTC)', 'Источник', 'Код страны', 'Страна', 'Доля (%)'], ([str(snapshots[g.content_snapshot_id].content_id), title(posts[snapshots[g.content_snapshot_id].content_id]), snapshots[g.content_snapshot_id].snapshot_at, snapshots[g.content_snapshot_id].source, g.country_code, g.country_name, g.percentage] for g in report.geography))
    out = BytesIO(); wb.save(out); return out.getvalue()


_FONT_LOCK = Lock()


def pdf_report(report):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, PageBreak
    from reportlab.graphics.shapes import Drawing, Rect, String
    with _FONT_LOCK:
        if 'Dash' not in pdfmetrics.getRegisteredFontNames():
            folder = Path(__file__).parent / 'fonts'
            pdfmetrics.registerFont(TTFont('Dash', str(folder / 'DejaVuSans.ttf')))
            pdfmetrics.registerFont(TTFont('DashBold', str(folder / 'DejaVuSans-Bold.ttf')))
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=42, bottomMargin=42, title='DASH - Аналитика аккаунтов', author='DASH Social Analytics')
    ink, green, muted = colors.HexColor('#1d2937'), colors.HexColor('#067462'), colors.HexColor('#566472')
    normal = ParagraphStyle('body', fontName='Dash', fontSize=9, leading=14, textColor=ink, spaceAfter=8, splitLongWords=True)
    heading = ParagraphStyle('heading', parent=normal, fontName='DashBold', fontSize=15, leading=20, spaceBefore=18, spaceAfter=10)
    small = ParagraphStyle('small', parent=normal, fontSize=8, leading=12, textColor=muted)
    def p(text, style=normal): return Paragraph(escape(str(text)), style)
    story = [p('DASH / SOCIAL ANALYTICS', small), p('Аналитика аккаунтов', ParagraphStyle('title', parent=heading, fontSize=25, leading=31, spaceBefore=6)), p(f'{report.period}  |  {report.zone.key}', normal), p(f'Сформирован {report.end.astimezone(report.zone):%d.%m.%Y в %H:%M}', small)]
    summary = report_summary(report)
    if not report.accounts:
        story += [Spacer(1, 16), p('Аккаунты ещё не добавлены')]
    for index, group in enumerate(summary['groups']):
        if index: story.append(PageBreak())
        section = [p(f'{group["platform"]} / @{group["username"]}', heading), p(f'Новых публикаций за период: {group["new_posts"]}')]
        if not group['measurements']: section.append(p('Подписчики: нет измерений за период', small))
        for row in group['measurements']:
            section.append(p(f'{row["source"]}: {number(row["followers"])} подписчиков / {row["at"]}'))
            section.append(p('Изменение: нужен второй замер' if row['change'] is None else f'Изменение между замерами: {row["change"]:+d} / {row["between"]}', small))
        story.append(KeepTogether(section))
        for block in group['tops']:
            posts = block['posts']
            # Draw bars with exact labels; never imply a weekly count from lifetime data.
            chart = Drawing(510, len(posts) * 31 + 12)
            maximum = max((post['views'] for post in posts), default=0)
            for i, post in enumerate(posts):
                y = (len(posts) - i - 1) * 31 + 8
                name = post['title'].replace('\n', ' ')
                while pdfmetrics.stringWidth(name, 'Dash', 8) > 230: name = name[:-2] + '…'
                chart.add(String(0, y + 4, name, fontName='Dash', fontSize=8, fillColor=ink))
                chart.add(Rect(244, y, 180, 15, fillColor=colors.HexColor('#e6f4ee'), strokeColor=None, rx=3, ry=3))
                width = 180 * post['views'] / maximum if maximum else 0
                if width > 0:
                    chart.add(Rect(244, y, width, 15, fillColor=green, strokeColor=None, rx=min(3, width / 2), ry=3))
                chart.add(String(434, y + 4, number(post['views']), fontName='DashBold', fontSize=8, fillColor=ink))
            story.append(KeepTogether([Spacer(1, 6), p(f'Просмотры публикаций, накопленные / {block["source"]}', small), chart]))
            rows = [[p('Публикация', small), p('Лайки', small), p('Комментарии', small), p('Репосты', small)]] + [[p(post['title'][:100]), p(number(post['likes'])), p(number(post['comments'])), p(number(post['shares']))] for post in posts]
            table = Table(rows, colWidths=[278, 65, 93, 79], repeatRows=1)
            table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f0f4f7')), ('LINEBELOW', (0, 0), (-1, -1), .4, colors.HexColor('#dfe5ec')), ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 1)]))
            story.append(table)
        if not group['tops']: story.append(p('Нет накопленных измерений просмотров за выбранный период', small))
    story += [Spacer(1, 18), p('Как читать отчёт', heading), p('Период выбирает даты сохранённых измерений. Накопленные просмотры публикаций не равны просмотрам за неделю. Изменение подписчиков рассчитано между указанными замерами, а не за полные календарные дни. Сводка предпочитает API; дополнительные источники и измерения сохранены отдельно в Excel. «Нет данных» означает отсутствие измерения; ноль - известное значение. Полная история измерений за период доступна в Excel.', small)]
    def footer(canvas, document):
        canvas.setFont('Dash', 8); canvas.setFillColor(muted)
        canvas.drawString(40, 22, 'DASH / отчёт по сохранённым данным')
        canvas.drawRightString(A4[0] - 40, 22, str(document.page))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
